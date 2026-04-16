# ||Shri Ganeshay Namaha||

import tkinter
from tkinter import filedialog, messagebox
from openpyxl import load_workbook
import pandas as pd
import os
import re
from pathlib import Path
from datetime import datetime

# ================= Global Variables =================
selected_df = None
folder_name = None
file_path = None
# ================= FILE READ =================
def open_file(): 
    global selected_df
    
    file_path = filedialog.askopenfilename(
        filetypes=[("Excel files", "*.xlsx")]
    )
    
    if not file_path:
        return

    label.config(text=f"Selected:\n{file_path}")
    
    try:
        wb = load_workbook(file_path)
        sheet = wb.active

        data = list(sheet.values)

        # Remove top empty rows
        while data and all(cell is None or str(cell).strip() == "" for cell in data[0]):
            data.pop(0)

        columns = data[0]
        rows = data[1:]

        df = pd.DataFrame(rows, columns=columns)

        # Clean data
        df = df.replace(r'^\s*$', None, regex=True)
        df = df.dropna(how="all")
        df = df.dropna(axis=1, how="all")

        # Skip repeated header row
        if str(df.iloc[0]["Row Labels"]).lower() == "row labels":
            df = df.iloc[1:]

        df = df.reset_index(drop=True)
        df.index = df.index + 1

        selected_df = df
        dat_button.config(state="normal")

        text_box.delete("1.0", tkinter.END)
        text_box.insert(tkinter.END, df.to_string())

    except Exception as e:
        text_box.delete("1.0", tkinter.END)
        text_box.insert(tkinter.END, f"Error: {e}")

# ================= CREATE DAT FILES =================
def create_dat_file():
    global selected_df, folder_name, file_path
    
    if selected_df is None:
        messagebox.showerror("Error", "No file selected")
        return

    df = selected_df
    now = datetime.now()
    cell_value = str(df.iloc[2, 1])
    substring = result = cell_value[2:4] + cell_value[6:8]
     
    folder_name = f"BA_Files_{substring}_{now.strftime('%H%M')}"
    
    if not os.path.exists(folder_name):
        os.makedirs(folder_name)

    try:
        for i in range(len(df)):
            row = df.iloc[i]

            Billing_Unit = row["Row Labels"]
            Date1 = str(row["Date"])

            Energy_Chrg = row.get("Sum of Energy Charge", 0) or 0
            Service_Chrg = row.get("Sum of ServiceCharge", 0) or 0
            CGST = row.get("Sum of CGST", 0) or 0
            SGST = row.get("Sum of SGST", 0) or 0

            # Tag
            tag = str(Billing_Unit).zfill(4) + "EV" + Date1.zfill(8) + "EVC"

            # File name
            fileName = f"BT_{Billing_Unit}_{Date1[0:2]}_{Date1[6:8]}{Date1[2:4]}_EV.DAT"
            file_path = os.path.join(folder_name, fileName)

            # Lines
            lin1 = tag + "EV01+" + str('{:.2f}'.format(Energy_Chrg)).zfill(15)
            lin2 = tag + "EV02+" + str('{:.2f}'.format(Service_Chrg)).zfill(15)
            lin3 = tag + "DA70+" + str('{:.2f}'.format(CGST)).zfill(15)
            lin4 = tag + "DB70+" + str('{:.2f}'.format(SGST)).zfill(15)

            data = lin1 + '\n' + lin2 + '\n' + lin3 + '\n' + lin4 + '\n'

            # Write file
            with open(file_path, 'w') as f:
                f.write(data)

        # Success message
        messagebox.showinfo("Success", f"DAT files created in:\n{folder_name}")

        # Open folder automatically
        try:
            os.startfile(os.path.abspath(folder_name))  # Windows
        except AttributeError:
            import subprocess
            subprocess.Popen(["open", os.path.abspath(folder_name)])  # Mac
            # subprocess.Popen(["xdg-open", os.path.abspath(folder_name)])  # Linux

        # Enable Summary button
        summary_button.config(state="normal")

    except Exception as e:
        messagebox.showerror("Error", str(e))

# ================= CREATE SUMMARY TXT =================
def create_summary():
    global folder_name, selected_df

    if selected_df is None:
        messagebox.showerror("Error", "No data available")
        return

    if not folder_name or not os.path.exists(folder_name):
        messagebox.showerror("Error", "DAT folder not found")
        return

    df = selected_df

    # ✅ Data size validation
    if df.shape[0] <= 2 or df.shape[1] <= 1:
        messagebox.showerror("Error", "Not enough data in file")
        return

    try:
        # ✅ SAFE extraction
        cell_value = str(df.iloc[2, 1])

        if len(cell_value) >= 8:
            substring = cell_value[2:4] + cell_value[6:8]
        else:
            substring = "NA"

        summary_file = os.path.join(
            folder_name,
            f"BA_Files_{substring}_All.txt"
        )

        # ✅ File writing
        with open(summary_file, 'w') as outfile:
            line_count = 0  # ✅ counter
            for file in os.listdir(folder_name):
                if file.endswith(".DAT"):
                    file_path = os.path.join(folder_name, file)
                    with open(file_path, 'r') as f:
                        for line in f:   # ✅ line is defined here
                            if line.strip():   # ✅ removes empty lines
                                line_count += 1   # ✅ increment count
                                outfile.write(line)

        # ✅ Show count in message
        #Use integer division: //
        messagebox.showinfo(
        "Success",
         f"All DAT files copied to:\n{summary_file}\n\nTotal records: {line_count // 4}"
        )

    except Exception as e:
        messagebox.showerror("Error", str(e))
# ================= UI =================
window = tkinter.Tk()
window.title("BA File Generator : (Excel to DAT)")
window.geometry("900x550")
window.configure(bg="#1e1e2f")

title = tkinter.Label(
    window,
    text="📊 BA File Generator : (Excel to DAT)",
    font=("Segoe UI", 20, "bold"),
    fg="#ffffff",
    bg="#1e1e2f"
)
title.pack(pady=10)

label = tkinter.Label(
    window,
    text="Select Excel file to generate DAT",
    font=("Segoe UI", 11),
    fg="#cfcfcf",
    bg="#1e1e2f",
    bd=2,
    relief="groove",
    padx=10,
    pady=10,
    wraplength=800,
    justify="center"
)
label.pack(pady=10)

btn_frame = tkinter.Frame(window, bg="#1e1e2f")
btn_frame.pack(pady=10)

select_button = tkinter.Button(
    btn_frame,
    text="Select Excel File",
    command=open_file,
    font=("Segoe UI", 12, "bold"),
    bg="#4a90e2",
    fg="white",
    padx=20,
    pady=10,
    bd=0,
    cursor="hand2"
)
select_button.pack(side="left", padx=10)

dat_button = tkinter.Button(
    btn_frame,
    text="Create DAT Files",
    command=create_dat_file,
    font=("Segoe UI", 12, "bold"),
    bg="#28a745",
    fg="white",
    padx=20,
    pady=10,
    bd=0,
    cursor="hand2",
    state="disabled"
)
dat_button.pack(side="left", padx=10)

summary_button = tkinter.Button(
    btn_frame,
    text="Create Summary",
    command=create_summary,
    font=("Segoe UI", 12, "bold"),
    bg="#ff9900",
    fg="white",
    padx=20,
    pady=10,
    bd=0,
    cursor="hand2",
    state="disabled"  # will enable after DAT files created
)
summary_button.pack(side="left", padx=10)

text_frame = tkinter.Frame(window)
text_frame.pack(fill="both", expand=True, padx=10, pady=10)

scrollbar = tkinter.Scrollbar(text_frame)
scrollbar.pack(side="right", fill="y")

text_box = tkinter.Text(
    text_frame,
    yscrollcommand=scrollbar.set,
    wrap="none",
    font=("Consolas", 10)
)
text_box.pack(side="left", fill="both", expand=True)

scrollbar.config(command=text_box.yview)

window.mainloop()
 